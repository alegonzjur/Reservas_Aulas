"""Pruebas de la aplicación. Se ejecutan contra PostgreSQL real (hace falta btree_gist):

    python manage.py test
"""
import csv
import os
import re
from datetime import date, time, timedelta
from io import BytesIO, StringIO
from smtplib import SMTPException
from unittest import mock

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core import mail
from django.core.management import call_command
from django.core.management.base import CommandError
from django.db import IntegrityError, transaction
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from . import emails
from .calendario import (
    ANCHO_EJE_PX, ANCHO_MIN_DIA_PX, Bloque, construir_semana, lunes_de, lunes_por_defecto,
    pendientes_para_revisar, repartir_en_carriles, texto_rango,
)
from .forms import PeticionForm
from .models import Aula, Curso, EventoReserva, Material, Reserva, SerieReserva
from .permissions import GRUPO_ADMIN, GRUPO_PROFESORADO
from .services import (
    ConflictoReserva,
    cambiar_estado,
    cambiar_estado_serie,
    crear_peticion,
    crear_serie,
    fechas_de_serie,
)

User = get_user_model()
Estado = Reserva.Estado
Tipo = EventoReserva.Tipo

CLAVE = "clave-solo-para-tests-123"
HOY = timezone.localdate()
# Lunes de la semana que viene: siempre futuro y día lectivo.
LUNES = HOY + timedelta(days=7 - HOY.weekday())


def t(texto: str) -> time:
    return time.fromisoformat(texto)


class BaseTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.g_admin = Group.objects.create(name=GRUPO_ADMIN)
        cls.g_prof = Group.objects.create(name=GRUPO_PROFESORADO)
        cls.mantenimiento = User.objects.create_user(
            "mantenimiento", password=CLAVE, first_name="Mantenimiento", is_staff=True
        )
        cls.mantenimiento.groups.add(cls.g_admin)
        cls.informatica = User.objects.create_user(
            "informatica", password=CLAVE, first_name="Informática", is_staff=True, is_superuser=True
        )
        cls.informatica.groups.add(cls.g_admin)
        cls.profesor = User.objects.create_user("profesor", password=CLAVE, first_name="Profesor")
        cls.profesor.groups.add(cls.g_prof)
        cls.sin_rol = User.objects.create_user("sinrol", password=CLAVE)
        cls.a1 = Aula.objects.create(nombre="Aula 1")
        cls.a2 = Aula.objects.create(nombre="Aula 2")
        cls.curso = Curso.objects.create(nombre="1ºA")
        cls.curso2 = Curso.objects.create(nombre="Grupo 2")
        # Los dos materiales de partida los siembra la migración 0004.
        cls.mat_maniquies = Material.objects.get(nombre="Maniquíes")
        cls.mat_otro = Material.objects.get(nombre="Material distinto al habitual")

    def crear(self, aula=None, ini="10:00", fin="11:00", estado=Estado.PENDIENTE, dia=None,
              profesor="Ana López", notas="", **extra):
        if estado == Estado.ACEPTADA:
            extra.setdefault("gestionada_por", self.mantenimiento)
            extra.setdefault("gestionada_en", timezone.now())
        return Reserva.objects.create(
            aula=aula or self.a1, fecha=dia or LUNES, hora_inicio=t(ini), hora_fin=t(fin),
            profesor=profesor, notas=notas, estado=estado, creada_por=self.profesor, **extra,
        )

    def datos_form(self, **cambios):
        datos = {
            "aula": self.a1.pk, "fecha": LUNES.isoformat(), "hora_inicio": "10:00",
            "hora_fin": "11:00", "titulo": "Anatomía I", "curso": self.curso.pk,
            "profesor": "Ana López", "notas": "",
        }
        datos.update(cambios)
        return datos


# =============================================================================== Base de datos
class RestriccionBaseDeDatosTests(BaseTest):
    """La garantía anti-doble-reserva vive en PostgreSQL, no solo en Python."""

    def test_solape_parcial_de_aceptadas_es_rechazado(self):
        self.crear(ini="10:00", fin="11:00", estado=Estado.ACEPTADA)
        with self.assertRaises(IntegrityError), transaction.atomic():
            self.crear(ini="10:30", fin="11:30", estado=Estado.ACEPTADA)

    def test_una_dentro_de_otra_es_rechazada(self):
        self.crear(ini="09:00", fin="12:00", estado=Estado.ACEPTADA)
        with self.assertRaises(IntegrityError), transaction.atomic():
            self.crear(ini="10:00", fin="11:00", estado=Estado.ACEPTADA)

    def test_franjas_contiguas_no_se_solapan(self):
        self.crear(ini="10:00", fin="11:00", estado=Estado.ACEPTADA)
        self.crear(ini="11:00", fin="12:00", estado=Estado.ACEPTADA)
        self.crear(ini="09:00", fin="10:00", estado=Estado.ACEPTADA)
        self.assertEqual(Reserva.objects.filter(estado=Estado.ACEPTADA).count(), 3)

    def test_misma_franja_en_otra_aula_o_otro_dia_es_valida(self):
        self.crear(estado=Estado.ACEPTADA)
        self.crear(aula=self.a2, estado=Estado.ACEPTADA)
        self.crear(dia=LUNES + timedelta(days=1), estado=Estado.ACEPTADA)
        self.assertEqual(Reserva.objects.count(), 3)

    def test_pendientes_y_rechazadas_pueden_solaparse_con_aceptadas(self):
        self.crear(estado=Estado.ACEPTADA)
        self.crear(estado=Estado.PENDIENTE)
        self.crear(estado=Estado.PENDIENTE)
        self.crear(estado=Estado.RECHAZADA)
        self.assertEqual(Reserva.objects.count(), 4)

    def test_fin_anterior_o_igual_al_inicio_es_rechazado(self):
        with self.assertRaises(IntegrityError), transaction.atomic():
            self.crear(ini="12:00", fin="11:00")
        with self.assertRaises(IntegrityError), transaction.atomic():
            self.crear(ini="12:00", fin="12:00")

    def test_migraciones_sincronizadas_con_los_modelos(self):
        # Falla si alguien cambia un modelo y olvida generar la migración.
        salida = StringIO()
        try:
            call_command("makemigrations", check=True, dry_run=True, stdout=salida)
        except SystemExit:
            self.fail(f"Hay cambios de modelo sin migración:\n{salida.getvalue()}")

    def test_validacion_de_modelo_usa_el_mensaje_de_la_restriccion(self):
        self.crear(estado=Estado.ACEPTADA)
        otra = Reserva(
            aula=self.a1, fecha=LUNES, hora_inicio=t("10:30"), hora_fin=t("11:30"),
            profesor="X", estado=Estado.ACEPTADA, creada_por=self.profesor,
        )
        from django.core.exceptions import ValidationError
        with self.assertRaises(ValidationError) as cm:
            otra.full_clean()
        self.assertIn("se solapa", str(cm.exception))


# =============================================================================== Servicios
@override_settings(RESERVAS_EMAIL_DESTINO=["mantenimiento@centro.test"], DEFAULT_FROM_EMAIL="reservas@centro.test")
class ServiciosTests(BaseTest):
    def nueva(self, **kw):
        return Reserva(
            aula=kw.get("aula", self.a1), fecha=LUNES, hora_inicio=t(kw.get("ini", "10:00")),
            hora_fin=t(kw.get("fin", "11:00")), profesor="Ana López", notas="Proyector",
        )

    def test_peticion_queda_pendiente_audita_y_avisa_por_correo(self):
        with self.captureOnCommitCallbacks(execute=True):
            r = crear_peticion(self.nueva(), self.profesor)
        r.refresh_from_db()
        self.assertEqual(r.estado, Estado.PENDIENTE)
        self.assertEqual(r.creada_por, self.profesor)
        self.assertIsNone(r.gestionada_por)
        self.assertEqual([e.tipo for e in r.eventos.all()], [Tipo.CREADA])

        self.assertEqual(len(mail.outbox), 1)
        correo = mail.outbox[0]
        self.assertEqual(correo.to, ["mantenimiento@centro.test"])
        self.assertEqual(correo.from_email, "reservas@centro.test")
        self.assertIn("Aula 1", correo.subject)
        self.assertIn("Ana López", correo.body)
        self.assertIn("Proyector", correo.body)
        self.assertIn(reverse("panel"), correo.body)

    def test_el_texto_libre_no_va_en_el_asunto_del_correo(self):
        r = self.nueva()
        r.profesor = "Ana\nBcc: intruso@x.com"
        with self.captureOnCommitCallbacks(execute=True):
            crear_peticion(r, self.profesor)
        self.assertNotIn("intruso", mail.outbox[0].subject)
        self.assertNotIn("\n", mail.outbox[0].subject)

    def test_reserva_directa_de_admin_queda_aceptada_sin_correo(self):
        with self.captureOnCommitCallbacks(execute=True):
            r = crear_peticion(self.nueva(), self.mantenimiento, aceptar=True)
        r.refresh_from_db()
        self.assertEqual(r.estado, Estado.ACEPTADA)
        self.assertEqual(r.gestionada_por, self.mantenimiento)
        self.assertIsNotNone(r.gestionada_en)
        self.assertEqual([e.tipo for e in r.eventos.all()], [Tipo.CREADA, Tipo.ACEPTADA])
        self.assertEqual(mail.outbox, [])

    def test_reserva_directa_sobre_franja_ocupada_falla_y_no_guarda_nada(self):
        self.crear(estado=Estado.ACEPTADA)
        antes = Reserva.objects.count()
        with self.assertRaises(ConflictoReserva) as cm:
            crear_peticion(self.nueva(ini="10:30", fin="11:30"), self.mantenimiento, aceptar=True)
        self.assertEqual(len(cm.exception.conflictos), 1)
        self.assertEqual(Reserva.objects.count(), antes)
        self.assertEqual(EventoReserva.objects.count(), 0)

    def test_aceptar_registra_quien_y_cuando(self):
        r = self.crear()
        antes = timezone.now()
        r, cambiado = cambiar_estado(r, Estado.ACEPTADA, self.mantenimiento)
        self.assertTrue(cambiado)
        r.refresh_from_db()
        self.assertEqual(r.estado, Estado.ACEPTADA)
        self.assertEqual(r.gestionada_por, self.mantenimiento)
        self.assertGreaterEqual(r.gestionada_en, antes)
        evento = r.eventos.get()
        self.assertEqual((evento.tipo, evento.usuario), (Tipo.ACEPTADA, self.mantenimiento))

    def test_aceptar_dos_veces_no_duplica_auditoria(self):
        r = self.crear()
        cambiar_estado(r, Estado.ACEPTADA, self.mantenimiento)
        _, cambiado = cambiar_estado(r, Estado.ACEPTADA, self.informatica)
        self.assertFalse(cambiado)
        self.assertEqual(r.eventos.count(), 1)
        r.refresh_from_db()
        self.assertEqual(r.gestionada_por, self.mantenimiento)  # no se pisa

    def test_no_se_puede_aceptar_sobre_una_franja_ocupada(self):
        self.crear(estado=Estado.ACEPTADA)
        solapada = self.crear(ini="10:30", fin="11:30", profesor="Luis")
        with self.assertRaises(ConflictoReserva):
            cambiar_estado(solapada, Estado.ACEPTADA, self.mantenimiento)
        solapada.refresh_from_db()
        self.assertEqual(solapada.estado, Estado.PENDIENTE)
        self.assertEqual(solapada.eventos.count(), 0)

    def test_elegir_entre_dos_pendientes_bloquea_la_otra(self):
        a = self.crear(profesor="Ana")
        b = self.crear(profesor="Luis")
        cambiar_estado(a, Estado.ACEPTADA, self.mantenimiento)
        with self.assertRaises(ConflictoReserva):
            cambiar_estado(b, Estado.ACEPTADA, self.mantenimiento)
        cambiar_estado(b, Estado.RECHAZADA, self.mantenimiento, "Ya está ocupada")
        b.refresh_from_db()
        self.assertEqual(b.estado, Estado.RECHAZADA)
        self.assertEqual(b.motivo_rechazo, "Ya está ocupada")

    def test_si_la_comprobacion_previa_falla_la_base_de_datos_lo_impide(self):
        """Simula una carrera: la comprobación en Python no ve el conflicto, pero la
        restricción de PostgreSQL sí, y se traduce a ConflictoReserva (no a un 500)."""
        self.crear(estado=Estado.ACEPTADA)
        solapada = self.crear(ini="10:30", fin="11:30")
        with mock.patch.object(Reserva, "aceptadas_solapadas", return_value=Reserva.objects.none()):
            with self.assertRaises(ConflictoReserva):
                cambiar_estado(solapada, Estado.ACEPTADA, self.mantenimiento)
        solapada.refresh_from_db()
        self.assertEqual(solapada.estado, Estado.PENDIENTE)

    def test_rechazar_y_reabrir(self):
        r = self.crear()
        cambiar_estado(r, Estado.RECHAZADA, self.mantenimiento, "  No hay proyector  ")
        r.refresh_from_db()
        self.assertEqual(r.motivo_rechazo, "No hay proyector")
        self.assertEqual(r.gestionada_por, self.mantenimiento)

        cambiar_estado(r, Estado.PENDIENTE, self.informatica)
        r.refresh_from_db()
        self.assertEqual(r.estado, Estado.PENDIENTE)
        self.assertIsNone(r.gestionada_por)
        self.assertIsNone(r.gestionada_en)
        self.assertEqual(r.motivo_rechazo, "")
        self.assertEqual([e.tipo for e in r.eventos.all()], [Tipo.RECHAZADA, Tipo.REABIERTA])

    def test_cancelar_una_aceptada_libera_la_franja(self):
        r = self.crear(estado=Estado.ACEPTADA)
        cambiar_estado(r, Estado.RECHAZADA, self.mantenimiento, "Cancelada")
        nueva = self.crear(ini="10:30", fin="11:30")
        cambiar_estado(nueva, Estado.ACEPTADA, self.mantenimiento)
        nueva.refresh_from_db()
        self.assertEqual(nueva.estado, Estado.ACEPTADA)

    def test_un_fallo_de_smtp_no_rompe_la_peticion(self):
        with mock.patch.object(emails, "send_mail", side_effect=SMTPException("caído")):
            with self.assertLogs("reservas.emails", level="ERROR"):
                with self.captureOnCommitCallbacks(execute=True):
                    r = crear_peticion(self.nueva(), self.profesor)
        self.assertTrue(Reserva.objects.filter(pk=r.pk).exists())

    @override_settings(RESERVAS_EMAIL_DESTINO=[])
    def test_sin_destinatarios_configurados_no_envia_ni_falla(self):
        r = self.crear()
        with self.assertLogs("reservas.emails", level="WARNING"):
            self.assertFalse(emails.notificar_peticion_pendiente(r.pk))
        self.assertEqual(mail.outbox, [])


# =============================================================================== Formulario
class PeticionFormTests(BaseTest):
    def form(self, es_admin=False, **cambios):
        return PeticionForm(self.datos_form(**cambios), es_admin=es_admin)

    def test_peticion_valida(self):
        f = self.form()
        self.assertTrue(f.is_valid(), f.errors)

    def test_fin_debe_ser_posterior_al_inicio(self):
        f = self.form(hora_inicio="11:00", hora_fin="10:00")
        self.assertFalse(f.is_valid())
        self.assertIn("posterior", str(f.errors))

    def test_limites_de_horario(self):
        self.assertTrue(self.form(hora_inicio="08:00", hora_fin="09:00").is_valid())
        self.assertTrue(self.form(hora_inicio="20:00", hora_fin="21:00").is_valid())
        self.assertFalse(self.form(hora_inicio="07:55", hora_fin="09:00").is_valid())
        self.assertFalse(self.form(hora_inicio="20:00", hora_fin="21:05").is_valid())

    def test_solo_dias_lectivos(self):
        sabado = LUNES + timedelta(days=5)
        f = self.form(fecha=sabado.isoformat())
        self.assertFalse(f.is_valid())
        self.assertIn("lunes a viernes", str(f.errors))

    def test_no_acepta_fechas_pasadas(self):
        ayer = HOY - timedelta(days=1)
        f = self.form(fecha=ayer.isoformat())
        self.assertFalse(f.is_valid())
        self.assertIn("ya han pasado", str(f.errors))

    def test_las_horas_deben_ser_multiplos_del_paso(self):
        f = self.form(hora_inicio="10:03")
        self.assertFalse(f.is_valid())
        self.assertIn("múltiplos de 5", str(f.errors))
        self.assertTrue(self.form(hora_inicio="10:05", hora_fin="14:40").is_valid())

    def test_bloquea_si_la_franja_esta_ocupada_por_una_aceptada(self):
        self.crear(estado=Estado.ACEPTADA, ini="10:30", fin="11:30")
        f = self.form()
        self.assertFalse(f.is_valid())
        self.assertIn("ya tiene una reserva aceptada", str(f.non_field_errors()))
        self.assertIn("10:30-11:30", str(f.non_field_errors()))

    def test_no_revela_quien_tiene_reservada_el_aula(self):
        self.crear(estado=Estado.ACEPTADA, profesor="Persona Secreta")
        f = self.form()
        self.assertNotIn("Secreta", str(f.errors))

    def test_varias_pendientes_en_la_misma_franja_si_se_permiten(self):
        self.crear(estado=Estado.PENDIENTE)
        self.assertTrue(self.form().is_valid())

    def test_contigua_a_una_aceptada_es_valida(self):
        self.crear(estado=Estado.ACEPTADA, ini="09:00", fin="10:00")
        self.assertTrue(self.form().is_valid())

    def test_un_aula_inactiva_no_se_puede_elegir(self):
        self.a2.activa = False
        self.a2.save()
        f = self.form(aula=self.a2.pk)
        self.assertFalse(f.is_valid())
        self.assertIn("aula", f.errors)

    def test_aula_obligatoria_y_profesor_obligatorio(self):
        f = self.form(aula="", profesor="")
        self.assertFalse(f.is_valid())
        self.assertIn("aula", f.errors)
        self.assertIn("profesor", f.errors)

    def test_titulo_y_curso_son_obligatorios(self):
        f = self.form(titulo="", curso="")
        self.assertFalse(f.is_valid())
        self.assertIn("titulo", f.errors)
        self.assertIn("curso", f.errors)

    def test_un_titulo_de_solo_espacios_no_vale(self):
        f = self.form(titulo="    ")
        self.assertFalse(f.is_valid())
        self.assertIn("titulo", f.errors)

    def test_el_titulo_se_limpia_y_tiene_maximo(self):
        f = self.form(titulo="  Anatomía I  ")
        self.assertTrue(f.is_valid(), f.errors)
        self.assertEqual(f.cleaned_data["titulo"], "Anatomía I")
        self.assertFalse(self.form(titulo="x" * 121).is_valid())

    def test_un_curso_inactivo_no_se_puede_elegir(self):
        self.curso2.activo = False
        self.curso2.save()
        f = self.form(curso=self.curso2.pk)
        self.assertFalse(f.is_valid())
        self.assertIn("curso", f.errors)

    def test_el_material_es_opcional_y_admite_varios(self):
        self.assertTrue(self.form().is_valid())
        f = self.form(materiales=[self.mat_maniquies.pk, self.mat_otro.pk])
        self.assertTrue(f.is_valid(), f.errors)
        self.assertEqual(set(f.cleaned_data["materiales"]), {self.mat_maniquies, self.mat_otro})

    def test_un_material_inactivo_no_se_puede_elegir(self):
        self.mat_otro.activo = False
        self.mat_otro.save()
        f = self.form(materiales=[self.mat_otro.pk])
        self.assertFalse(f.is_valid())
        self.assertIn("materiales", f.errors)

    def test_la_casilla_aceptar_directamente_solo_existe_para_admins(self):
        self.assertNotIn("aceptar_directamente", self.form(es_admin=False).fields)
        self.assertIn("aceptar_directamente", self.form(es_admin=True).fields)


# =============================================================================== Calendario
class CalendarioTests(BaseTest):
    def carriles(self, *franjas):
        reservas = [
            Reserva(pk=i + 1, aula=self.a1, fecha=LUNES, hora_inicio=t(a), hora_fin=t(b), profesor="x")
            for i, (a, b) in enumerate(franjas)
        ]
        return {r.pk: (c, n) for r, c, n in repartir_en_carriles(reservas)}

    def test_sin_solapes_cada_una_ocupa_todo_el_ancho(self):
        self.assertEqual(self.carriles(("08:00", "09:00"), ("10:00", "11:00")), {1: (0, 1), 2: (0, 1)})

    def test_contiguas_no_cuentan_como_solape(self):
        self.assertEqual(self.carriles(("08:00", "09:00"), ("09:00", "10:00")), {1: (0, 1), 2: (0, 1)})

    def test_dos_solapadas_se_colocan_lado_a_lado(self):
        self.assertEqual(self.carriles(("10:00", "11:00"), ("10:30", "11:30")), {1: (0, 2), 2: (1, 2)})

    def test_cadena_de_solapes_reutiliza_carriles(self):
        # A(8-10) B(9-11) C(10-12): C puede ir en el carril de A porque A ya terminó.
        res = self.carriles(("08:00", "10:00"), ("09:00", "11:00"), ("10:00", "12:00"))
        self.assertEqual(res, {1: (0, 2), 2: (1, 2), 3: (0, 2)})

    def test_tres_simultaneas_usan_tres_carriles(self):
        res = self.carriles(("10:00", "11:00"), ("10:00", "11:00"), ("10:00", "11:00"))
        self.assertEqual(sorted(res.values()), [(0, 3), (1, 3), (2, 3)])

    def test_el_estilo_css_usa_punto_decimal(self):
        b = Bloque(reserva=None, top=15.384615, alto=7.692307, carril=1, carriles=3)
        self.assertEqual(b.estilo, "--top:15.385%;--alto:7.692%;--izq:33.333%;--ancho:33.333%")

    def test_semana_muestra_pendientes_y_aceptadas_pero_no_rechazadas(self):
        self.crear(estado=Estado.ACEPTADA, ini="10:00", fin="11:00")
        self.crear(aula=self.a2, estado=Estado.PENDIENTE, ini="12:00", fin="13:00", dia=LUNES + timedelta(days=1))
        self.crear(estado=Estado.RECHAZADA, ini="15:00", fin="16:00")
        self.crear(dia=LUNES + timedelta(days=7))  # otra semana
        semana = construir_semana(LUNES)

        self.assertEqual(len(semana["dias"]), 5)
        self.assertEqual([d["nombre"] for d in semana["dias"]],
                         ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes"])
        self.assertEqual(len(semana["dias"][0]["bloques"]), 1)
        self.assertEqual(len(semana["dias"][1]["bloques"]), 1)
        self.assertEqual(sum(len(d["bloques"]) for d in semana["dias"]), 2)
        self.assertTrue(semana["hay_reservas"])

    def test_geometria_de_un_bloque(self):
        self.crear(ini="10:00", fin="11:00", estado=Estado.ACEPTADA)
        bloque = construir_semana(LUNES)["dias"][0]["bloques"][0]
        self.assertAlmostEqual(bloque.top, 120 / 780 * 100, places=4)   # 2 h tras las 08:00
        self.assertAlmostEqual(bloque.alto, 60 / 780 * 100, places=4)

    def test_primera_y_ultima_franja_caben_en_la_cuadricula(self):
        self.crear(ini="08:00", fin="09:00")
        self.crear(aula=self.a2, ini="20:00", fin="21:00")
        primero, segundo = sorted(construir_semana(LUNES)["dias"][0]["bloques"], key=lambda b: b.top)
        self.assertEqual(primero.top, 0)
        self.assertAlmostEqual(segundo.top + segundo.alto, 100, places=4)

    def test_filtro_por_aula(self):
        self.crear(aula=self.a1)
        self.crear(aula=self.a2)
        self.assertEqual(len(construir_semana(LUNES, self.a2.pk)["dias"][0]["bloques"]), 1)

    def test_semana_vacia(self):
        self.assertFalse(construir_semana(LUNES)["hay_reservas"])

    def test_texto_rango_mismo_mes_cruce_de_mes_y_de_año(self):
        self.assertEqual(texto_rango(date(2026, 10, 19), date(2026, 10, 23)), "Del 19 al 23 de octubre de 2026")
        self.assertEqual(
            texto_rango(date(2026, 10, 26), date(2026, 10, 30)), "Del 26 al 30 de octubre de 2026"
        )
        self.assertEqual(
            texto_rango(date(2026, 11, 30), date(2026, 12, 4)), "Del 30 de noviembre al 4 de diciembre de 2026"
        )
        self.assertEqual(
            texto_rango(date(2026, 12, 28), date(2027, 1, 1)),
            "Del 28 de diciembre de 2026 al 1 de enero de 2027",
        )

    def test_columnas_crecen_con_los_carriles_simultaneos(self):
        vacia = construir_semana(LUNES)
        self.assertEqual(vacia["ancho_minimo"], ANCHO_EJE_PX + 5 * ANCHO_MIN_DIA_PX)
        self.assertTrue(vacia["columnas_css"].startswith(f"{ANCHO_EJE_PX}px "))
        for ini, fin in (("10:00", "11:00"), ("10:00", "11:00"), ("10:00", "11:00"), ("10:00", "11:00")):
            self.crear(ini=ini, fin=fin)
        llena = construir_semana(LUNES)
        # 4 carriles x 100px = 400px para el lunes; el resto de días se queda en el mínimo.
        self.assertIn("minmax(400px, 1fr)", llena["columnas_css"])
        self.assertEqual(llena["ancho_minimo"], ANCHO_EJE_PX + 400 + 4 * ANCHO_MIN_DIA_PX)

    def test_las_horas_del_eje_van_de_las_8_a_las_21(self):
        horas = construir_semana(LUNES)["horas"]
        self.assertEqual((horas[0]["texto"], horas[-1]["texto"], len(horas)), ("08:00", "21:00", 14))

    def test_semana_por_defecto_salta_al_lunes_siguiente_en_fin_de_semana(self):
        miercoles = LUNES + timedelta(days=2)
        sabado, domingo = LUNES + timedelta(days=5), LUNES + timedelta(days=6)
        self.assertEqual(lunes_por_defecto(miercoles), LUNES)
        self.assertEqual(lunes_por_defecto(sabado), LUNES + timedelta(days=7))
        self.assertEqual(lunes_por_defecto(domingo), LUNES + timedelta(days=7))
        self.assertEqual(lunes_de(domingo), LUNES)

    def test_avisos_en_la_lista_de_pendientes(self):
        self.crear(estado=Estado.ACEPTADA, ini="10:00", fin="11:00")
        choca = self.crear(ini="10:30", fin="11:30", profesor="Choca")
        rival1 = self.crear(aula=self.a2, ini="12:00", fin="13:00", profesor="R1")
        rival2 = self.crear(aula=self.a2, ini="12:30", fin="13:30", profesor="R2")
        sola = self.crear(aula=self.a2, ini="16:00", fin="17:00", profesor="Sola")
        vieja = self.crear(dia=HOY - timedelta(days=3), profesor="Vieja")

        pendientes, total = pendientes_para_revisar()
        por_nombre = {p.profesor: p for p in pendientes}
        self.assertEqual(total, 5)
        self.assertTrue(por_nombre["Choca"].conflicto)
        self.assertEqual(por_nombre["R1"].competidoras, 1)
        self.assertEqual(por_nombre["R2"].competidoras, 1)
        self.assertFalse(por_nombre["R1"].conflicto)
        self.assertEqual((por_nombre["Sola"].conflicto, por_nombre["Sola"].competidoras), (False, 0))
        self.assertTrue(por_nombre["Vieja"].caducada)
        self.assertFalse(por_nombre["Sola"].caducada)
        self.assertEqual(pendientes[0].profesor, "Vieja")  # ordenadas por fecha


# =============================================================================== Permisos y vistas
class PermisosTests(BaseTest):
    def test_anonimo_va_al_login(self):
        for nombre in ("panel", "reservar", "listado", "exportar"):
            r = self.client.get(reverse(nombre))
            self.assertRedirects(r, f"{reverse('login')}?next={reverse(nombre)}", fetch_redirect_response=False)

    def test_profesorado_no_accede_a_la_gestion(self):
        self.client.force_login(self.profesor)
        r = self.crear()
        for url in (reverse("listado"), reverse("exportar")):
            self.assertEqual(self.client.get(url).status_code, 403, url)
        for accion in ("aceptar", "rechazar", "reabrir"):
            resp = self.client.post(reverse("reserva_accion", args=[r.pk]), {"accion": accion})
            self.assertEqual(resp.status_code, 403, accion)
        r.refresh_from_db()
        self.assertEqual(r.estado, Estado.PENDIENTE)
        self.assertEqual(r.eventos.count(), 0)

    def test_usuario_sin_grupo_no_puede_ni_reservar(self):
        self.client.force_login(self.sin_rol)
        self.assertEqual(self.client.get(reverse("reservar")).status_code, 403)
        self.assertEqual(self.client.get(reverse("panel")).status_code, 403)

    def test_la_pagina_403_es_la_de_la_aplicacion(self):
        self.client.force_login(self.profesor)
        r = self.client.get(reverse("listado"))
        self.assertContains(r, "No tienes permiso", status_code=403)

    def test_el_inicio_lleva_a_cada_rol_a_su_pantalla(self):
        self.client.force_login(self.profesor)
        self.assertRedirects(self.client.get(reverse("inicio")), reverse("reservar"), fetch_redirect_response=False)
        self.client.force_login(self.mantenimiento)
        self.assertRedirects(self.client.get(reverse("inicio")), reverse("panel"), fetch_redirect_response=False)
        self.client.force_login(self.informatica)
        self.assertRedirects(self.client.get(reverse("inicio")), reverse("panel"), fetch_redirect_response=False)

    def test_login_real_redirige_segun_rol(self):
        r = self.client.post(reverse("login"), {"username": "mantenimiento", "password": CLAVE}, follow=True)
        self.assertEqual(r.redirect_chain[-1][0], reverse("panel"))
        self.client.logout()
        r = self.client.post(reverse("login"), {"username": "profesor", "password": CLAVE}, follow=True)
        self.assertEqual(r.redirect_chain[-1][0], reverse("reservar"))

    def test_login_con_clave_incorrecta(self):
        r = self.client.post(reverse("login"), {"username": "profesor", "password": "mal"})
        self.assertContains(r, "no son correctos")

    def test_logout_solo_por_post(self):
        self.client.force_login(self.profesor)
        self.assertEqual(self.client.get(reverse("logout")).status_code, 405)
        self.assertRedirects(self.client.post(reverse("logout")), reverse("login"), fetch_redirect_response=False)

    def test_csrf_se_exige_en_las_acciones(self):
        cliente = Client(enforce_csrf_checks=True)
        cliente.force_login(self.mantenimiento)
        r = self.crear()
        resp = cliente.post(reverse("reserva_accion", args=[r.pk]), {"accion": "aceptar"})
        self.assertEqual(resp.status_code, 403)
        r.refresh_from_db()
        self.assertEqual(r.estado, Estado.PENDIENTE)


@override_settings(RESERVAS_EMAIL_DESTINO=["mantenimiento@centro.test"])
class FormularioVistaTests(BaseTest):
    def test_el_profesorado_ve_el_formulario_sin_la_casilla_de_admin(self):
        self.client.force_login(self.profesor)
        r = self.client.get(reverse("reservar"))
        self.assertContains(r, "Pedir un aula")
        self.assertContains(r, "data-buscable")
        self.assertNotContains(r, "aceptar_directamente")

    def test_el_admin_ve_la_casilla_aceptar_directamente(self):
        self.client.force_login(self.mantenimiento)
        self.assertContains(self.client.get(reverse("reservar")), "aceptar_directamente")

    def test_profesor_envia_peticion_queda_pendiente_y_se_avisa(self):
        self.client.force_login(self.profesor)
        with self.captureOnCommitCallbacks(execute=True):
            r = self.client.post(reverse("reservar"), self.datos_form(profesor="Otra Persona"), follow=True)
        self.assertRedirects(r, reverse("reservar"))
        reserva = Reserva.objects.get()
        self.assertEqual(
            (reserva.estado, reserva.profesor, reserva.creada_por), (Estado.PENDIENTE, "Otra Persona", self.profesor)
        )
        self.assertContains(r, "Petición enviada")
        self.assertEqual(len(mail.outbox), 1)

    def test_la_peticion_guarda_titulo_curso_y_material_y_el_correo_los_incluye(self):
        self.client.force_login(self.profesor)
        datos = self.datos_form(
            titulo="Taller de RCP", curso=self.curso2.pk,
            materiales=[self.mat_maniquies.pk, self.mat_otro.pk],
        )
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(reverse("reservar"), datos)
        reserva = Reserva.objects.get()
        self.assertEqual((reserva.titulo, reserva.curso), ("Taller de RCP", self.curso2))
        self.assertEqual(set(reserva.materiales.all()), {self.mat_maniquies, self.mat_otro})
        cuerpo = mail.outbox[0].body
        self.assertIn("Título:    Taller de RCP", cuerpo)
        self.assertIn("Curso:     Grupo 2", cuerpo)
        self.assertIn("Maniquíes", cuerpo)
        self.assertNotIn("Taller de RCP", mail.outbox[0].subject)  # el texto libre no va al asunto

    def test_sin_material_marcado_se_guarda_vacio(self):
        self.client.force_login(self.profesor)
        self.client.post(reverse("reservar"), self.datos_form())
        self.assertEqual(Reserva.objects.get().materiales.count(), 0)

    def test_el_formulario_sin_titulo_o_curso_no_guarda_nada(self):
        self.client.force_login(self.profesor)
        for campo in ("titulo", "curso"):
            r = self.client.post(reverse("reservar"), self.datos_form(**{campo: ""}))
            self.assertEqual(r.status_code, 200)
            self.assertContains(r, "obligatorio")
        self.assertEqual(Reserva.objects.count(), 0)
        self.assertEqual(mail.outbox, [])

    def test_el_formulario_pinta_los_campos_nuevos(self):
        self.client.force_login(self.profesor)
        r = self.client.get(reverse("reservar"))
        self.assertContains(r, 'name="titulo"')
        self.assertContains(r, "1ºA")
        self.assertContains(r, 'name="materiales"')
        self.assertContains(r, "Maniquíes")
        self.assertContains(r, "Material distinto al habitual")

    def test_el_material_se_guarda_tambien_en_reserva_directa_de_admin(self):
        self.client.force_login(self.mantenimiento)
        self.client.post(
            reverse("reservar"),
            self.datos_form(aceptar_directamente="on", materiales=[self.mat_maniquies.pk]),
        )
        r = Reserva.objects.get()
        self.assertEqual(r.estado, Estado.ACEPTADA)
        self.assertEqual(list(r.materiales.all()), [self.mat_maniquies])

    def test_el_profesorado_no_puede_colarse_una_reserva_aceptada(self):
        self.client.force_login(self.profesor)
        self.client.post(reverse("reservar"), self.datos_form(aceptar_directamente="on"))
        self.assertEqual(Reserva.objects.get().estado, Estado.PENDIENTE)

    def test_admin_reserva_directa_queda_aceptada(self):
        self.client.force_login(self.mantenimiento)
        self.client.post(reverse("reservar"), self.datos_form(aceptar_directamente="on"))
        r = Reserva.objects.get()
        self.assertEqual((r.estado, r.gestionada_por), (Estado.ACEPTADA, self.mantenimiento))

    def test_admin_sin_marcar_la_casilla_la_deja_pendiente(self):
        self.client.force_login(self.mantenimiento)
        self.client.post(reverse("reservar"), self.datos_form())
        self.assertEqual(Reserva.objects.get().estado, Estado.PENDIENTE)

    def test_peticion_sobre_franja_ocupada_se_bloquea_con_aviso(self):
        self.crear(estado=Estado.ACEPTADA)
        self.client.force_login(self.profesor)
        r = self.client.post(reverse("reservar"), self.datos_form(hora_inicio="10:30", hora_fin="11:30"))
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "ya tiene una reserva aceptada")
        self.assertEqual(Reserva.objects.count(), 1)

    def test_errores_de_validacion_se_muestran_y_no_se_guarda(self):
        self.client.force_login(self.profesor)
        r = self.client.post(reverse("reservar"), self.datos_form(hora_inicio="12:00", hora_fin="11:00"))
        self.assertContains(r, "posterior")
        self.assertEqual(Reserva.objects.count(), 0)
        self.assertEqual(mail.outbox, [])

    def test_aviso_en_vivo_de_disponibilidad(self):
        self.crear(estado=Estado.ACEPTADA, ini="10:00", fin="11:00")
        self.client.force_login(self.profesor)
        url = reverse("disponibilidad")
        base = {"aula": self.a1.pk, "fecha": LUNES.isoformat(), "hora_inicio": "10:30", "hora_fin": "11:30"}
        self.assertContains(self.client.get(url, base), "Aula ocupada")
        libre = dict(base, hora_inicio="11:00", hora_fin="12:00")
        self.assertContains(self.client.get(url, libre), "Franja libre")
        incompleta = {"aula": self.a1.pk}
        self.assertNotContains(self.client.get(url, incompleta), "Franja")
        self.assertEqual(self.client.get(url, incompleta).status_code, 200)
        invertida = dict(base, hora_inicio="12:00", hora_fin="11:00")
        self.assertNotContains(self.client.get(url, invertida), "Franja")


class CalendarioProfesoradoTests(BaseTest):
    """El profesorado ve el calendario (para saber si hay reservas) pero solo en lectura."""

    HTMX = {"HTTP_HX_REQUEST": "true"}

    def setUp(self):
        self.client.force_login(self.profesor)
        self.aceptada = self.crear(
            estado=Estado.ACEPTADA, profesor="Aceptada Ana", titulo="Anatomía I", curso=self.curso,
            notas="NOTA INTERNA", ini="09:00", fin="10:00",
        )
        self.aceptada.materiales.set([self.mat_maniquies])
        self.pendiente = self.crear(
            aula=self.a2, profesor="Pendiente Pau", titulo="Taller de RCP", ini="12:00", fin="13:00",
            dia=LUNES + timedelta(days=1),
        )
        self.rechazada = self.crear(
            profesor="Rechazada Rita", estado=Estado.RECHAZADA, ini="15:00", fin="16:00", titulo="Cancelada",
        )

    def panel(self, **parametros):
        return self.client.get(reverse("panel"), {"semana": LUNES.isoformat(), **parametros})

    def test_ve_solo_las_reservas_aceptadas(self):
        r = self.panel()
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Aceptada Ana")
        self.assertContains(r, "bloque-aceptada")
        # Ni las pendientes ni las rechazadas existen para el profesorado.
        for oculto in ("Pendiente Pau", "Taller de RCP", "Rechazada Rita", "bloque-pendiente", "Pendiente de revisar"):
            self.assertNotContains(r, oculto)

    def test_el_fragmento_del_calendario_tampoco_filtra_pendientes(self):
        r = self.client.get(reverse("panel_contenido"), {"semana": LUNES.isoformat()})
        self.assertContains(r, "Aceptada Ana")
        for oculto in ("Pendiente Pau", "bloque-pendiente", "Pendiente de revisar"):
            self.assertNotContains(r, oculto)

    def test_el_titulo_aparece_en_los_bloques(self):
        r = self.panel()
        self.assertContains(r, "bloque-titulo")
        self.assertContains(r, "Anatomía I")

    def test_una_semana_con_solo_pendientes_se_ve_vacia(self):
        semana = LUNES + timedelta(days=14)
        self.crear(profesor="Solo Pendiente", dia=semana)
        r = self.client.get(reverse("panel"), {"semana": semana.isoformat()})
        self.assertContains(r, "No hay reservas esta semana")
        self.assertNotContains(r, "Solo Pendiente")

    def test_no_ve_la_lista_de_pendientes_ni_controles_de_gestion(self):
        for r in (self.panel(), self.client.get(reverse("panel_contenido"), {"semana": LUNES.isoformat()})):
            self.assertEqual(r.status_code, 200)
            self.assertNotContains(r, "Pendientes de revisar")
            self.assertNotContains(r, "hx-swap-oob")
            self.assertNotContains(r, "/accion/")
            self.assertNotContains(r, "reservasActualizadas")

    def test_puede_navegar_entre_semanas_y_filtrar_por_aula(self):
        r = self.client.get(reverse("panel_contenido"), {"semana": LUNES.isoformat(), "aula": self.a2.pk})
        self.assertNotContains(r, "Aceptada Ana")  # es del aula 1
        r = self.client.get(reverse("panel_contenido"), {"semana": LUNES.isoformat(), "aula": self.a1.pk})
        self.assertContains(r, "Aceptada Ana")
        siguiente = self.client.get(reverse("panel_contenido"), {"semana": (LUNES + timedelta(days=7)).isoformat()})
        self.assertEqual(siguiente.status_code, 200)

    def test_la_ficha_es_de_solo_lectura_y_no_filtra_datos_internos(self):
        r = self.client.get(reverse("reserva_detalle", args=[self.aceptada.pk]), **self.HTMX)
        self.assertEqual(r.status_code, 200)
        for visible in ("Aula 1", "Anatomía I", "1ºA", "Aceptada Ana", "09:00 a 10:00"):
            self.assertContains(r, visible)
        for oculto in ("NOTA INTERNA", "Maniquíes", "Historial", "Aceptar", "Rechazar", "Cancelar", "<form"):
            self.assertNotContains(r, oculto)

    def test_las_pendientes_y_rechazadas_no_existen_para_el_profesorado(self):
        for reserva in (self.pendiente, self.rechazada):
            r = self.client.get(reverse("reserva_detalle", args=[reserva.pk]), **self.HTMX)
            self.assertEqual(r.status_code, 404, reserva.estado)
            self.assertNotContains(r, "Pendiente Pau", status_code=404)

    def test_una_pendiente_aceptada_despues_pasa_a_verse(self):
        cambiar_estado(self.pendiente, Estado.ACEPTADA, self.mantenimiento)
        self.assertContains(self.panel(), "Pendiente Pau")
        r = self.client.get(reverse("reserva_detalle", args=[self.pendiente.pk]), **self.HTMX)
        self.assertEqual(r.status_code, 200)

    def test_la_ficha_completa_sin_htmx_solo_es_para_administradores(self):
        r = self.client.get(reverse("reserva_detalle", args=[self.aceptada.pk]))
        self.assertRedirects(r, reverse("panel"), fetch_redirect_response=False)

    def test_no_puede_gestionar_aunque_lo_intente_a_mano(self):
        for accion in ("aceptar", "rechazar", "reabrir"):
            r = self.client.post(reverse("reserva_accion", args=[self.pendiente.pk]), {"accion": accion}, **self.HTMX)
            self.assertEqual(r.status_code, 403)
        self.pendiente.refresh_from_db()
        self.assertEqual(self.pendiente.estado, Estado.PENDIENTE)

    def test_el_menu_del_profesorado_tiene_calendario_pero_no_gestion(self):
        r = self.client.get(reverse("reservar"))
        self.assertContains(r, f'href="{reverse("panel")}"')
        self.assertContains(r, "Pedir un aula")
        for oculto in ("Exportar datos", "Aulas y ajustes", f'href="{reverse("listado")}"'):
            self.assertNotContains(r, oculto)

    def test_sin_rol_sigue_sin_ver_nada(self):
        self.client.force_login(self.sin_rol)
        self.assertEqual(self.panel().status_code, 403)
        self.assertEqual(self.client.get(reverse("reserva_detalle", args=[self.aceptada.pk]), **self.HTMX).status_code, 403)

    def test_el_administrador_sigue_viendo_ficha_completa_y_pendientes(self):
        self.client.force_login(self.mantenimiento)
        r = self.panel()
        self.assertContains(r, "Pendientes de revisar")
        self.assertContains(r, "/accion/")
        ficha = self.client.get(reverse("reserva_detalle", args=[self.aceptada.pk]), **self.HTMX)
        for visible in ("NOTA INTERNA", "Maniquíes", "Historial", "Anatomía I", "1ºA"):
            self.assertContains(ficha, visible)


class PanelVistaTests(BaseTest):
    def setUp(self):
        self.client.force_login(self.mantenimiento)

    def test_panel_muestra_calendario_y_pendientes(self):
        self.crear(estado=Estado.ACEPTADA, profesor="Aceptada Ana", dia=LUNES)
        self.crear(aula=self.a2, profesor="Pendiente Pau", dia=LUNES + timedelta(days=1))
        r = self.client.get(reverse("panel"), {"semana": LUNES.isoformat()})
        self.assertContains(r, "Aceptada Ana")
        self.assertContains(r, "Pendiente Pau")
        self.assertContains(r, "bloque-aceptada")
        self.assertContains(r, "bloque-pendiente")
        self.assertContains(r, "Lunes")
        self.assertContains(r, LUNES.strftime("%d/%m/%Y"))
        self.assertContains(r, "Pendientes de revisar")
        self.assertContains(r, "08:00")
        self.assertContains(r, "21:00")

    def test_panel_pinta_rechazadas_en_ningun_sitio_del_calendario(self):
        self.crear(estado=Estado.RECHAZADA, profesor="Rechazada Rosa")
        self.assertNotContains(self.client.get(reverse("panel"), {"semana": LUNES.isoformat()}), "Rechazada Rosa")

    def test_parametros_invalidos_no_rompen_el_panel(self):
        for params in ({"semana": "basura"}, {"aula": "x"}, {"semana": "2026-13-45", "aula": "9999"}):
            self.assertEqual(self.client.get(reverse("panel"), params).status_code, 200, params)

    def test_cualquier_dia_de_la_semana_lleva_a_su_lunes(self):
        r = self.client.get(reverse("panel"), {"semana": (LUNES + timedelta(days=3)).isoformat()})
        self.assertEqual(r.context["lunes"], LUNES)

    def test_el_css_no_recibe_decimales_con_coma(self):
        self.crear(ini="10:00", fin="11:00", estado=Estado.ACEPTADA)
        html = self.client.get(reverse("panel"), {"semana": LUNES.isoformat()}).content.decode()
        estilos = re.findall(r'style="(--top:[^"]+)"', html)
        self.assertTrue(estilos)
        for estilo in estilos:
            self.assertNotIn(",", estilo)

    def test_fragmento_htmx_incluye_pendientes_fuera_de_banda_y_actualiza_la_url(self):
        self.crear()
        r = self.client.get(reverse("panel_contenido"), {"semana": LUNES.isoformat()}, headers={"HX-Request": "true"})
        self.assertContains(r, 'hx-swap-oob="true"')
        self.assertEqual(r["HX-Push-Url"], f"{reverse('panel')}?semana={LUNES.isoformat()}")
        self.assertNotContains(r, "<html")

    def test_el_refresco_automatico_no_toca_el_historial_del_navegador(self):
        r = self.client.get(reverse("panel_contenido"), {"semana": LUNES.isoformat(), "refresco": "1"})
        self.assertNotIn("HX-Push-Url", r)

    def test_la_pagina_completa_no_duplica_la_lista_de_pendientes(self):
        r = self.client.get(reverse("panel"), {"semana": LUNES.isoformat()})
        self.assertNotContains(r, "hx-swap-oob")
        self.assertEqual(r.content.decode().count('id="pendientes"'), 1)

    def test_el_filtro_de_aula_se_conserva_en_los_enlaces_de_semana(self):
        r = self.client.get(reverse("panel"), {"semana": LUNES.isoformat(), "aula": self.a2.pk})
        self.assertContains(r, f"aula={self.a2.pk}")
        self.assertContains(r, f'<option value="{self.a2.pk}" selected>')

    def test_lista_de_pendientes_oculta_aceptar_si_hay_conflicto(self):
        self.crear(estado=Estado.ACEPTADA)
        self.crear(ini="10:30", fin="11:30", profesor="Choca")
        r = self.client.get(reverse("panel"), {"semana": LUNES.isoformat()})
        self.assertContains(r, "Aula ya ocupada")
        self.assertNotContains(r, 'value="aceptar"')


class AccionesVistaTests(BaseTest):
    def setUp(self):
        self.client.force_login(self.mantenimiento)

    def accion(self, reserva, accion, htmx=True, **extra):
        cabeceras = {"HX-Request": "true"} if htmx else {}
        return self.client.post(
            reverse("reserva_accion", args=[reserva.pk]), dict(accion=accion, **extra), headers=cabeceras
        )

    def test_detalle_como_fragmento_y_como_pagina(self):
        r = self.crear(notas="Traer proyector")
        frag = self.client.get(reverse("reserva_detalle", args=[r.pk]), headers={"HX-Request": "true"})
        self.assertContains(frag, "Traer proyector")
        self.assertNotContains(frag, "<html")
        pagina = self.client.get(reverse("reserva_detalle", args=[r.pk]))
        self.assertContains(pagina, "<html")
        self.assertContains(pagina, "Traer proyector")
        self.assertEqual(self.client.get(reverse("reserva_detalle", args=[99999])).status_code, 404)

    def test_aceptar_por_htmx_cambia_estado_y_pide_recargar_el_panel(self):
        r = self.crear()
        resp = self.accion(r, "aceptar")
        self.assertEqual(resp["HX-Trigger"], "reservasActualizadas")
        self.assertContains(resp, "Reserva aceptada")
        self.assertContains(resp, "Aceptada por")
        self.assertContains(resp, "Mantenimiento")  # quién
        r.refresh_from_db()
        self.assertEqual((r.estado, r.gestionada_por), (Estado.ACEPTADA, self.mantenimiento))

    def test_aceptar_sobre_ocupada_muestra_error_y_no_cambia_nada(self):
        self.crear(estado=Estado.ACEPTADA, profesor="Ocupa Olga")
        r = self.crear(ini="10:30", fin="11:30", profesor="Choca")
        resp = self.accion(r, "aceptar")
        self.assertContains(resp, "No se puede aceptar")
        self.assertContains(resp, "Ocupa Olga")
        self.assertContains(resp, "disabled")
        r.refresh_from_db()
        self.assertEqual(r.estado, Estado.PENDIENTE)

    def test_rechazar_con_motivo(self):
        r = self.crear()
        resp = self.accion(r, "rechazar", motivo="Aula en obras")
        self.assertContains(resp, "Aula en obras")
        r.refresh_from_db()
        self.assertEqual((r.estado, r.motivo_rechazo), (Estado.RECHAZADA, "Aula en obras"))

    def test_reabrir_una_rechazada(self):
        r = self.crear(estado=Estado.RECHAZADA)
        self.accion(r, "reabrir")
        r.refresh_from_db()
        self.assertEqual(r.estado, Estado.PENDIENTE)

    def test_doble_clic_no_falla_ni_duplica(self):
        r = self.crear()
        self.accion(r, "aceptar")
        resp = self.accion(r, "aceptar")
        self.assertContains(resp, "ya estaba aceptada")
        self.assertEqual(r.eventos.count(), 1)

    def test_el_historial_muestra_el_rastro(self):
        r = self.crear()
        self.accion(r, "aceptar")
        resp = self.client.get(reverse("reserva_detalle", args=[r.pk]), headers={"HX-Request": "true"})
        self.assertContains(resp, "Historial")
        self.assertContains(resp, "Aceptada por Mantenimiento")

    def test_accion_invalida_y_metodo_incorrecto(self):
        r = self.crear()
        self.assertEqual(self.accion(r, "borrar").status_code, 400)
        self.assertEqual(self.client.get(reverse("reserva_accion", args=[r.pk])).status_code, 405)

    def test_sin_javascript_redirige_al_detalle_con_mensaje(self):
        r = self.crear()
        resp = self.accion(r, "aceptar", htmx=False)
        # fetch_redirect_response=False: seguir la redirección aquí consumiría el mensaje.
        self.assertRedirects(resp, reverse("reserva_detalle", args=[r.pk]), fetch_redirect_response=False)
        seguida = self.client.get(resp.url)
        self.assertContains(seguida, "Reserva aceptada")

    def test_dos_administradores_distintos_quedan_auditados_por_separado(self):
        r = self.crear()
        self.accion(r, "aceptar")
        self.client.force_login(self.informatica)
        self.accion(r, "rechazar", motivo="Cambio de planes")
        usuarios = [(e.tipo, e.usuario.username) for e in r.eventos.all()]
        self.assertEqual(usuarios, [(Tipo.ACEPTADA, "mantenimiento"), (Tipo.RECHAZADA, "informatica")])


class ListadoVistaTests(BaseTest):
    def setUp(self):
        self.client.force_login(self.mantenimiento)
        self.crear(profesor="Ana", estado=Estado.ACEPTADA, notas="proyector")
        self.crear(profesor="Berta", estado=Estado.PENDIENTE, aula=self.a2, dia=LUNES + timedelta(days=1))
        self.crear(profesor="Carla", estado=Estado.RECHAZADA, dia=LUNES + timedelta(days=2))

    def test_sin_filtros_lista_todo(self):
        r = self.client.get(reverse("listado"))
        self.assertEqual(len(r.context["pagina"].object_list), 3)

    def test_filtros(self):
        def nombres(**p):
            return {x.profesor for x in self.client.get(reverse("listado"), p).context["pagina"].object_list}
        self.assertEqual(nombres(estado="pendiente"), {"Berta"})
        self.assertEqual(nombres(aula=self.a2.pk), {"Berta"})
        self.assertEqual(nombres(q="proyector"), {"Ana"})
        self.assertEqual(nombres(q="carl"), {"Carla"})
        self.assertEqual(nombres(desde=(LUNES + timedelta(days=1)).isoformat()), {"Berta", "Carla"})
        self.assertEqual(nombres(hasta=LUNES.isoformat()), {"Ana"})

    def test_rango_de_fechas_incoherente_muestra_error(self):
        r = self.client.get(reverse("listado"), {"desde": "2026-12-31", "hasta": "2026-01-01"})
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "no puede ser posterior")

    def test_paginacion(self):
        for i in range(35):
            self.crear(profesor=f"Extra {i}", ini="08:00", fin="09:00", dia=LUNES + timedelta(days=3))
        r = self.client.get(reverse("listado"))
        self.assertEqual(len(r.context["pagina"].object_list), 30)
        self.assertContains(r, "Página 1 de 2")
        self.assertEqual(len(self.client.get(reverse("listado"), {"pagina": 2}).context["pagina"].object_list), 8)


# =============================================================================== Exportación
class ExportacionTests(BaseTest):
    def setUp(self):
        self.client.force_login(self.mantenimiento)
        self.ana = self.crear(
            profesor="Ana López", notas="Con ñ y tildes: á é", estado=Estado.ACEPTADA,
            titulo="Anatomía I", curso=self.curso,
        )
        self.ana.materiales.set([self.mat_otro, self.mat_maniquies])
        self.berta = self.crear(profesor="Berta", aula=self.a2, ini="12:00", fin="13:30")
        crear_evento = EventoReserva.objects.create
        crear_evento(reserva=self.ana, tipo=Tipo.CREADA, usuario=self.profesor)
        crear_evento(reserva=self.ana, tipo=Tipo.ACEPTADA, usuario=self.mantenimiento)
        crear_evento(reserva=self.berta, tipo=Tipo.CREADA, usuario=self.profesor)

    def descargar(self, formato, **filtros):
        return self.client.get(reverse("exportar"), dict(formato=formato, **filtros))

    def filas_csv(self, respuesta, delimitador=","):
        texto = respuesta.content.decode("utf-8-sig")
        return list(csv.reader(StringIO(texto), delimiter=delimitador))

    def test_csv_de_reservas(self):
        r = self.descargar("reservas_csv")
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r["Content-Type"].startswith("text/csv"))
        self.assertIn("attachment", r["Content-Disposition"])
        self.assertTrue(r.content.startswith(b"\xef\xbb\xbf"))  # BOM para Excel
        filas = self.filas_csv(r)
        cabecera = filas[0]
        self.assertEqual(
            cabecera[:11],
            ["id", "aula", "fecha", "hora_inicio", "hora_fin", "titulo", "curso", "profesor",
             "materiales", "notas", "estado"],
        )
        datos = {f[cabecera.index("profesor")]: dict(zip(cabecera, f)) for f in filas[1:]}
        ana = datos["Ana López"]
        self.assertEqual((ana["aula"], ana["fecha"], ana["hora_inicio"], ana["hora_fin"], ana["estado"]),
                         ("Aula 1", LUNES.isoformat(), "10:00", "11:00", "aceptada"))
        self.assertEqual(ana["notas"], "Con ñ y tildes: á é")
        self.assertEqual(ana["gestionada_por"], "mantenimiento")
        self.assertEqual((ana["titulo"], ana["curso"]), ("Anatomía I", "1ºA"))
        self.assertEqual(ana["materiales"], "Maniquíes | Material distinto al habitual")  # en el orden del catálogo
        self.assertEqual(datos["Berta"]["gestionada_por"], "")
        self.assertEqual((datos["Berta"]["titulo"], datos["Berta"]["curso"], datos["Berta"]["materiales"]), ("", "", ""))
        self.assertRegex(ana["creada_en"], r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$")

    def test_csv_con_punto_y_coma(self):
        r = self.descargar("reservas_csv", delimitador=";")
        filas = self.filas_csv(r, ";")
        self.assertEqual(filas[0][0], "id")
        self.assertEqual(len(filas), 3)

    def test_csv_neutraliza_formulas(self):
        self.crear(profesor='=HYPERLINK("http://malo")', notas="+cmd|' /C calc'!A0", ini="14:00", fin="15:00")
        self.crear(profesor="@SUM(1)", notas="-2+3", ini="15:00", fin="16:00")
        filas = self.filas_csv(self.descargar("reservas_csv"))
        celdas = {c for f in filas[1:] for c in f}
        self.assertIn("'=HYPERLINK(\"http://malo\")", celdas)
        self.assertIn("'+cmd|' /C calc'!A0", celdas)
        self.assertIn("'@SUM(1)", celdas)
        self.assertIn("'-2+3", celdas)
        # Ninguna celda del fichero puede empezar por un carácter que Excel interprete como fórmula.
        for f in filas[1:]:
            for c in f:
                self.assertFalse(c.startswith(("=", "+", "-", "@")), c)

    def test_excel_tiene_las_hojas_con_tipos_reales(self):
        r = self.descargar("xlsx")
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        libro = load_workbook(BytesIO(r.content))
        self.assertEqual(libro.sheetnames, ["Reservas", "Aulas", "Cursos", "Materiales", "Auditoría"])

        hoja = libro["Reservas"]
        cabecera = [c.value for c in hoja[1]]
        fila = dict(zip(cabecera, [c.value for c in hoja[2]]))
        self.assertEqual(fila["profesor"], "Ana López")
        self.assertEqual(fila["fecha"].date(), LUNES)           # fecha real, no texto
        self.assertEqual(fila["hora_inicio"], time(10, 0))      # hora real
        self.assertEqual(hoja.freeze_panes, "A2")
        self.assertEqual(libro["Aulas"].max_row, 3)             # cabecera + 2 aulas
        self.assertEqual(libro["Cursos"].max_row, 3)            # cabecera + 2 cursos
        self.assertEqual(libro["Materiales"].max_row, 3)        # cabecera + 2 materiales de partida
        self.assertEqual(libro["Auditoría"].max_row, 4)         # cabecera + 3 eventos

    def test_excel_no_ejecuta_formulas(self):
        self.crear(profesor="=SUM(1+1)", ini="14:00", fin="15:00")
        libro = load_workbook(BytesIO(self.descargar("xlsx").content))
        hoja = libro["Reservas"]
        col = [c.value for c in hoja[1]].index("profesor")
        celdas = [fila[col] for fila in hoja.iter_rows(min_row=2)]
        peligrosa = next(c for c in celdas if c.value == "=SUM(1+1)")
        self.assertEqual(peligrosa.data_type, "s")  # texto, no fórmula

    def test_los_filtros_afectan_a_reservas_y_auditoria_pero_no_a_aulas(self):
        r = self.descargar("xlsx", estado="aceptada")
        libro = load_workbook(BytesIO(r.content))
        self.assertEqual(libro["Reservas"].max_row, 2)
        self.assertEqual(libro["Auditoría"].max_row, 3)   # solo los 2 eventos de Ana
        self.assertEqual(libro["Aulas"].max_row, 3)

    def test_csv_de_aulas_y_auditoria(self):
        self.a2.activa = False
        self.a2.save()
        aulas = self.filas_csv(self.descargar("aulas_csv"))
        self.assertEqual(aulas[0], ["id", "nombre", "ubicacion", "capacidad", "activa"])
        self.assertEqual({(f[1], f[4]) for f in aulas[1:]}, {("Aula 1", "Sí"), ("Aula 2", "No")})

        auditoria = self.filas_csv(self.descargar("auditoria_csv"))
        self.assertEqual(auditoria[0], ["id_evento", "id_reserva", "fecha", "tipo", "usuario", "detalle"])
        self.assertEqual(len(auditoria) - 1, 3)
        self.assertIn(["aceptada", "mantenimiento"], [f[3:5] for f in auditoria[1:]])

    def test_exportacion_vacia_solo_lleva_cabecera(self):
        filas = self.filas_csv(self.descargar("reservas_csv", q="no-existe"))
        self.assertEqual(len(filas), 1)

    def test_formato_desconocido_y_filtros_invalidos(self):
        self.assertEqual(self.descargar("pdf").status_code, 400)
        r = self.descargar("xlsx", desde="2026-12-31", hasta="2026-01-01")
        self.assertEqual(r.status_code, 200)  # vuelve a la página con el error
        self.assertContains(r, "no puede ser posterior")

    def test_la_pagina_de_exportacion(self):
        r = self.client.get(reverse("exportar"))
        self.assertContains(r, "Descargar Excel")
        self.assertContains(r, "Auditoría en CSV")

    def test_solo_administradores_pueden_exportar(self):
        self.client.force_login(self.profesor)
        self.assertEqual(self.descargar("xlsx").status_code, 403)
        self.assertEqual(self.descargar("reservas_csv").status_code, 403)


# =============================================================================== Comando de usuarios
ENTORNO = {
    "RESERVAS_PASSWORD_MANTENIMIENTO": "Mant-Clave-Test-2026!",
    "RESERVAS_PASSWORD_INFORMATICA": "Info-Clave-Test-2026!",
    "RESERVAS_PASSWORD_PROFESOR": "profe26",
}


class ComandoUsuariosTests(TestCase):
    def ejecutar(self, *args, entorno=ENTORNO):
        with mock.patch.dict(os.environ, entorno):
            call_command("crear_usuarios", "--no-input", *args, stdout=StringIO(), stderr=StringIO())

    def test_crea_grupos_usuarios_y_permisos(self):
        self.ejecutar()
        mant, info, prof = (User.objects.get(username=u) for u in ("mantenimiento", "informatica", "profesor"))
        self.assertEqual([g.name for g in mant.groups.all()], [GRUPO_ADMIN])
        self.assertEqual([g.name for g in prof.groups.all()], [GRUPO_PROFESORADO])
        self.assertTrue(mant.is_staff and not mant.is_superuser)
        self.assertTrue(info.is_staff and info.is_superuser)
        self.assertFalse(prof.is_staff or prof.is_superuser)
        self.assertTrue(mant.check_password(ENTORNO["RESERVAS_PASSWORD_MANTENIMIENTO"]))
        self.assertTrue(prof.check_password("profe26"))
        permisos = set(Group.objects.get(name=GRUPO_ADMIN).permissions.values_list("codename", flat=True))
        self.assertIn("add_aula", permisos)
        self.assertIn("add_curso", permisos)
        self.assertIn("change_material", permisos)
        self.assertIn("change_reserva", permisos)
        self.assertNotIn("delete_reserva", permisos)
        self.assertNotIn("add_reserva", permisos)

    def test_es_idempotente_y_no_pisa_contrasenas(self):
        self.ejecutar()
        otro = dict(ENTORNO, RESERVAS_PASSWORD_PROFESOR="otra-clave-9")
        self.ejecutar(entorno=otro)
        self.assertEqual(User.objects.count(), 3)
        self.assertTrue(User.objects.get(username="profesor").check_password("profe26"))

    def test_restablecer_cambia_la_contrasena(self):
        self.ejecutar()
        otro = dict(ENTORNO, RESERVAS_PASSWORD_PROFESOR="otra-clave-9")
        self.ejecutar("--restablecer", entorno=otro)
        self.assertTrue(User.objects.get(username="profesor").check_password("otra-clave-9"))

    def test_falta_una_variable_de_entorno(self):
        entorno = {k: v for k, v in ENTORNO.items() if k != "RESERVAS_PASSWORD_PROFESOR"}
        with mock.patch.dict(os.environ, entorno, clear=False):
            os.environ.pop("RESERVAS_PASSWORD_PROFESOR", None)
            with self.assertRaisesMessage(CommandError, "RESERVAS_PASSWORD_PROFESOR"):
                call_command("crear_usuarios", "--no-input", stdout=StringIO())

    def test_las_cuentas_de_administracion_exigen_contrasena_robusta(self):
        debil = dict(ENTORNO, RESERVAS_PASSWORD_MANTENIMIENTO="12345678")
        with self.assertRaises(CommandError):
            self.ejecutar(entorno=debil)

    def test_la_clave_del_profesorado_exige_longitud_minima(self):
        corta = dict(ENTORNO, RESERVAS_PASSWORD_PROFESOR="abc")
        with self.assertRaisesMessage(CommandError, "al menos 6"):
            self.ejecutar(entorno=corta)


# =============================================================================== Admin de Django
class AdminTests(BaseTest):
    def test_mantenimiento_no_puede_borrar_ni_crear_reservas(self):
        r = self.crear()
        self.client.force_login(self.mantenimiento)
        # Sin permisos de modelo asignados al grupo en este test, comprobamos las reglas propias.
        from django.contrib import admin as dj_admin
        from .admin import ReservaAdmin
        ma = ReservaAdmin(Reserva, dj_admin.site)
        peticion = mock.Mock(user=self.mantenimiento)
        self.assertFalse(ma.has_add_permission(peticion))
        self.assertFalse(ma.has_delete_permission(peticion, r))
        self.assertTrue(ma.has_delete_permission(mock.Mock(user=self.informatica), r))

    def test_editar_en_admin_deja_rastro_y_no_deja_cambiar_el_estado(self):
        r = self.crear()
        self.client.force_login(self.informatica)
        url = reverse("admin:reservas_reserva_change", args=[r.pk])
        datos = {
            "aula": self.a1.pk, "fecha": LUNES.isoformat(), "hora_inicio": "10:00:00",
            "hora_fin": "11:00:00", "profesor": "Ana López", "notas": "corregido",
            "estado": "aceptada",  # el campo es de solo lectura: debe ignorarse
        }
        resp = self.client.post(url, datos)
        self.assertEqual(resp.status_code, 302, getattr(resp, "context", None) and resp.context["adminform"].errors)
        r.refresh_from_db()
        self.assertEqual(r.notas, "corregido")
        self.assertEqual(r.estado, Estado.PENDIENTE)
        evento = r.eventos.get()
        self.assertEqual((evento.tipo, evento.usuario), (Tipo.EDITADA, self.informatica))

    def test_el_historial_no_se_puede_editar_desde_admin(self):
        from django.contrib import admin as dj_admin
        from .admin import EventoReservaAdmin
        ma = EventoReservaAdmin(EventoReserva, dj_admin.site)
        peticion = mock.Mock(user=self.informatica)
        self.assertFalse(ma.has_add_permission(peticion))
        self.assertFalse(ma.has_change_permission(peticion))
        self.assertFalse(ma.has_delete_permission(peticion))


# ----------------------------------------------------------------------- Despliegue
import logging  # noqa: E402
import tempfile  # noqa: E402
from pathlib import Path  # noqa: E402
from types import SimpleNamespace  # noqa: E402

import run  # noqa: E402  (ejecutor del servidor, en la raíz del proyecto)


class EjecutorArgumentosTests(TestCase):
    def test_valores_por_defecto(self):
        a = run.leer_argumentos([], entorno={})
        self.assertEqual((a.host, a.port, a.threads, a.check), ("0.0.0.0", 8000, 8, False))

    def test_variables_de_entorno(self):
        a = run.leer_argumentos(
            [], entorno={"RESERVAS_HOST": "10.0.0.5", "RESERVAS_PORT": "8080", "RESERVAS_THREADS": "4"}
        )
        self.assertEqual((a.host, a.port, a.threads), ("10.0.0.5", 8080, 4))

    def test_la_linea_de_comandos_manda_sobre_el_entorno(self):
        a = run.leer_argumentos(["--port", "9000", "--check"], entorno={"RESERVAS_PORT": "8080"})
        self.assertEqual(a.port, 9000)
        self.assertTrue(a.check)

    def test_numero_no_valido_en_el_entorno(self):
        with self.assertRaisesMessage(run.ErrorDeArranque, "RESERVAS_PORT"):
            run.leer_argumentos([], entorno={"RESERVAS_PORT": "abc"})

    def test_rangos(self):
        for argv in (["--port", "0"], ["--port", "70000"], ["--threads", "0"], ["--threads", "65"]):
            with self.subTest(argv=argv), self.assertRaises(SystemExit) as cm:
                with mock.patch("sys.stderr", new=StringIO()):
                    run.leer_argumentos(argv, entorno={})
            self.assertEqual(cm.exception.code, 2)


class EjecutorComprobacionesTests(TestCase):
    def ajustes(self, raiz, **extra):
        origen = Path(raiz) / "static"
        origen.mkdir(exist_ok=True)
        base = dict(
            DEBUG=False,
            ALLOWED_HOSTS=["reservas.centro.local"],
            SITE_URL="http://reservas.centro.local:8000",
            EMAIL_BACKEND="django.core.mail.backends.smtp.EmailBackend",
            RESERVAS_EMAIL_DESTINO="mantenimiento@centro.local",
            STATIC_ROOT=str(Path(raiz) / "staticfiles"),
            STATICFILES_DIRS=[str(origen)],
        )
        base.update(extra)
        return SimpleNamespace(**base)

    def test_estaticos_ausentes_impiden_arrancar(self):
        with tempfile.TemporaryDirectory() as d:
            with self.assertRaisesMessage(run.ErrorDeArranque, "collectstatic"):
                run.comprobar_estaticos(self.ajustes(d))

    def test_estaticos_presentes(self):
        with tempfile.TemporaryDirectory() as d:
            css = Path(d) / "staticfiles" / "css"
            css.mkdir(parents=True)
            (css / "app.css").write_text("x")
            run.comprobar_estaticos(self.ajustes(d))

    def test_estaticos_no_se_exigen_con_debug(self):
        with tempfile.TemporaryDirectory() as d:
            run.comprobar_estaticos(self.ajustes(d, DEBUG=True))

    def test_configuracion_correcta_sin_avisos(self):
        with tempfile.TemporaryDirectory() as d:
            self.assertEqual(run.recopilar_avisos(self.ajustes(d), "0.0.0.0"), [])

    def test_avisos_de_configuracion_de_desarrollo(self):
        with tempfile.TemporaryDirectory() as d:
            ajustes = self.ajustes(
                d,
                DEBUG=True,
                ALLOWED_HOSTS=["localhost", "127.0.0.1"],
                RESERVAS_EMAIL_DESTINO="",
            )
            texto = "\n".join(run.recopilar_avisos(ajustes, "0.0.0.0"))
        self.assertIn("DJANGO_DEBUG", texto)
        self.assertIn("ALLOWED_HOSTS", texto)
        self.assertIn("RESERVAS_EMAIL_DESTINO", texto)

    def test_avisos_de_correo_y_url_solo_sin_debug(self):
        with tempfile.TemporaryDirectory() as d:
            ajustes = self.ajustes(
                d,
                SITE_URL="http://localhost:8000",
                EMAIL_BACKEND="django.core.mail.backends.console.EmailBackend",
            )
            texto = "\n".join(run.recopilar_avisos(ajustes, "0.0.0.0"))
        self.assertIn("SITE_URL", texto)
        self.assertIn("EMAIL_BACKEND", texto)

    def test_aviso_de_estaticos_desactualizados(self):
        with tempfile.TemporaryDirectory() as d:
            ajustes = self.ajustes(d)
            recogidos = Path(ajustes.STATIC_ROOT)
            recogidos.mkdir()
            (recogidos / "a.css").write_text("viejo")
            nuevo = Path(ajustes.STATICFILES_DIRS[0]) / "a.css"
            nuevo.write_text("nuevo")
            os.utime(nuevo, (os.path.getmtime(recogidos / "a.css") + 60,) * 2)
            texto = "\n".join(run.recopilar_avisos(ajustes, "0.0.0.0"))
        self.assertIn("collectstatic", texto)

    def test_base_de_datos_alcanzable(self):
        self.assertTrue(run.comprobar_base_de_datos())

    def test_base_de_datos_caida(self):
        from django.db import connection
        from django.db.utils import OperationalError

        with mock.patch.object(connection, "ensure_connection", side_effect=OperationalError("sin conexión")):
            with self.assertRaisesMessage(run.ErrorDeArranque, "PostgreSQL"):
                run.comprobar_base_de_datos()


class PlantillasDeErrorTests(TestCase):
    def test_404_personalizada(self):
        r = Client().get("/no/existe/")
        self.assertEqual(r.status_code, 404)
        self.assertContains(r, "No encontramos esa página", status_code=404)

    def test_500_es_autonoma(self):
        from django.template.loader import render_to_string

        html = render_to_string("500.html")
        self.assertIn("Ha ocurrido un error en el servidor", html)
        self.assertNotIn("{%", html)
        self.assertNotIn("/static/", html)  # no depende de estáticos


class LoggingTests(TestCase):
    def test_los_errores_de_django_llegan_a_un_manejador(self):
        # Con DEBUG=False un 500 solo se enviaría a ADMINS (vacío); debe quedar también en el log.
        self.assertTrue(logging.getLogger("django").handlers)
        self.assertTrue(logging.getLogger("reservas").handlers)


class CatalogoExportarImportarTests(TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.dir.cleanup)
        self.ruta = str(Path(self.dir.name) / "catalogo.json")

    def exportar(self):
        call_command("exportar_catalogo", self.ruta, stdout=StringIO())

    def importar(self, *extra):
        out = StringIO()
        call_command("importar_catalogo", self.ruta, *extra, stdout=out)
        return out.getvalue()

    def test_ida_y_vuelta_sin_chocar_con_los_materiales_iniciales(self):
        Aula.objects.create(nombre="Laboratorio 1", ubicacion="Planta 2", capacidad=30)
        Aula.objects.create(nombre="Aula vieja", activa=False)
        Curso.objects.create(nombre="1ºA")
        Material.objects.create(nombre="Camillas", orden=5)
        self.exportar()
        # «Servidor»: base vacía, salvo los materiales que crea la migración 0004.
        Aula.objects.all().delete()
        Curso.objects.all().delete()
        Material.objects.filter(nombre="Camillas").delete()
        salida = self.importar()
        self.assertIn("Catálogo importado", salida)
        lab = Aula.objects.get(nombre="Laboratorio 1")
        self.assertEqual((lab.ubicacion, lab.capacidad, lab.activa), ("Planta 2", 30, True))
        self.assertFalse(Aula.objects.get(nombre="Aula vieja").activa)
        self.assertTrue(Curso.objects.filter(nombre="1ºA").exists())
        self.assertEqual(Material.objects.get(nombre="Camillas").orden, 5)
        self.assertEqual(Material.objects.filter(nombre="Maniquíes").count(), 1)

    def test_es_idempotente_y_actualiza(self):
        Aula.objects.create(nombre="A1", capacidad=10)
        self.exportar()
        self.assertIn("0 nuevos, 0 actualizados, 1 sin cambios", self.importar())
        Aula.objects.filter(nombre="A1").update(capacidad=99)
        self.assertIn("0 nuevos, 1 actualizados", self.importar())
        self.assertEqual(Aula.objects.get(nombre="A1").capacidad, 10)
        self.assertEqual(Aula.objects.count(), 1)

    def test_no_borra_lo_que_no_esta_en_el_fichero(self):
        self.exportar()
        Aula.objects.create(nombre="Solo en el servidor")
        self.importar()
        self.assertTrue(Aula.objects.filter(nombre="Solo en el servidor").exists())

    def test_simulacro_no_guarda(self):
        Aula.objects.create(nombre="A1")
        self.exportar()
        Aula.objects.all().delete()
        salida = self.importar("--simulacro")
        self.assertIn("1 nuevos", salida)
        self.assertIn("Simulacro", salida)
        self.assertEqual(Aula.objects.count(), 0)

    def test_fichero_invalido_no_guarda_nada(self):
        Path(self.ruta).write_text(
            '{"version": 1, "aulas": [{"nombre": "Buena"}, {"nombre": "Mala", "capacidad": -3}]}',
            encoding="utf-8",
        )
        with self.assertRaises(CommandError):
            self.importar()
        self.assertFalse(Aula.objects.filter(nombre="Buena").exists())

    def test_ficheros_que_no_son_un_catalogo(self):
        for contenido in ("no es json", '{"version": 99}', "[]"):
            Path(self.ruta).write_text(contenido, encoding="utf-8")
            with self.subTest(contenido=contenido), self.assertRaises(CommandError):
                self.importar()
        with self.assertRaises(CommandError):
            call_command("importar_catalogo", str(Path(self.dir.name) / "no_existe.json"))


# ================================================================================ Series semanales
S1, S2, S3, S4 = (LUNES + timedelta(days=7 * i) for i in range(4))  # cuatro lunes seguidos


class SerieFormularioTests(BaseTest):
    def form(self, **cambios):
        return PeticionForm(self.datos_form(**cambios))

    def test_fechas_de_serie_incluye_ambos_extremos_y_salta_de_siete_en_siete(self):
        self.assertEqual(fechas_de_serie(S1, S4), [S1, S2, S3, S4])
        self.assertEqual(fechas_de_serie(S1, S4 - timedelta(days=1)), [S1, S2, S3])
        self.assertEqual(fechas_de_serie(S1, S1), [S1])

    def test_sin_repetir_es_una_reserva_normal(self):
        f = self.form()
        self.assertTrue(f.is_valid(), f.errors)
        self.assertIsNone(f.fechas_serie)

    def test_repetir_hasta_genera_las_fechas(self):
        f = self.form(repetir_hasta=S4.isoformat())
        self.assertTrue(f.is_valid(), f.errors)
        self.assertEqual(f.fechas_serie, [S1, S2, S3, S4])
        self.assertEqual(f.fechas_ocupadas, [])

    def test_hasta_una_fecha_que_no_es_lunes_llega_al_ultimo_lunes_anterior(self):
        f = self.form(repetir_hasta=(S3 + timedelta(days=4)).isoformat())  # viernes
        self.assertTrue(f.is_valid(), f.errors)
        self.assertEqual(f.fechas_serie, [S1, S2, S3])

    def test_la_fecha_final_debe_dejar_al_menos_una_semana(self):
        for hasta in (S1, S1 + timedelta(days=6)):
            with self.subTest(hasta=hasta):
                f = self.form(repetir_hasta=hasta.isoformat())
                self.assertFalse(f.is_valid())
                self.assertIn("repetir_hasta", f.errors)

    def test_limite_de_fechas_por_serie(self):
        con_limite = LUNES + timedelta(days=7 * 39)  # 40 fechas: permitido
        self.assertTrue(self.form(repetir_hasta=con_limite.isoformat()).is_valid())
        f = self.form(repetir_hasta=(con_limite + timedelta(days=7)).isoformat())  # 41
        self.assertFalse(f.is_valid())
        self.assertIn("máximo 40", f.errors["repetir_hasta"][0])

    def test_las_fechas_ocupadas_no_impiden_pedir_las_demas(self):
        self.crear(dia=S2, estado=Estado.ACEPTADA)
        f = self.form(repetir_hasta=S4.isoformat())
        self.assertTrue(f.is_valid(), f.errors)
        self.assertEqual(f.fechas_ocupadas, [S2])

    def test_si_todas_las_fechas_estan_ocupadas_se_bloquea(self):
        for dia in (S1, S2):
            self.crear(dia=dia, estado=Estado.ACEPTADA)
        f = self.form(repetir_hasta=S2.isoformat())
        self.assertFalse(f.is_valid())
        self.assertIn("todas las fechas", f.non_field_errors()[0])

    def test_una_pendiente_no_cuenta_como_ocupada(self):
        self.crear(dia=S2)  # pendiente
        self.assertEqual(self.form(repetir_hasta=S4.isoformat()).fechas_ocupadas, [])


class SerieServiciosTests(BaseTest):
    def plantilla(self, **extra):
        return Reserva(
            aula=self.a1, fecha=S1, hora_inicio=t("10:00"), hora_fin=t("12:20"),
            profesor="Ana López", titulo="Anatomía I", curso=self.curso, notas="Traer batas", **extra,
        )

    def test_crea_una_reserva_por_semana_con_los_mismos_datos(self):
        with self.captureOnCommitCallbacks(execute=True):
            res = crear_serie(self.plantilla(), self.profesor, hasta=S4, materiales=[self.mat_maniquies])
        reservas = list(res.serie.reservas.order_by("fecha"))
        self.assertEqual([r.fecha for r in reservas], [S1, S2, S3, S4])
        for r in reservas:
            self.assertEqual(
                (r.estado, r.aula, r.hora_inicio, r.hora_fin, r.titulo, r.curso, r.notas, r.creada_por),
                (Estado.PENDIENTE, self.a1, t("10:00"), t("12:20"), "Anatomía I", self.curso,
                 "Traer batas", self.profesor),
            )
            self.assertEqual(list(r.materiales.all()), [self.mat_maniquies])
        self.assertEqual(res.omitidas, [])

    def test_un_unico_correo_por_serie(self):
        with self.captureOnCommitCallbacks(execute=True):
            crear_serie(self.plantilla(), self.profesor, hasta=S4)
        self.assertEqual(len(mail.outbox), 1)
        correo = mail.outbox[0]
        self.assertIn("4 fechas", correo.subject)
        self.assertIn("10:00 - 12:20", correo.body)
        self.assertIn(f"{S1:%d/%m/%Y}", correo.body)
        self.assertIn(f"{S4:%d/%m/%Y}", correo.body)
        self.assertIn("Anatomía I", correo.body)
        self.assertNotIn("Anatomía I", correo.subject)

    def test_cada_fecha_tiene_su_historial(self):
        res = crear_serie(self.plantilla(), self.profesor, hasta=S2)
        for r in res.serie.reservas.all():
            evento = r.eventos.get()
            self.assertEqual(evento.tipo, Tipo.CREADA)
            self.assertIn(f"Serie {res.serie.pk}", evento.detalle)

    def test_reserva_directa_de_admin_acepta_todas_sin_correo(self):
        with self.captureOnCommitCallbacks(execute=True):
            res = crear_serie(self.plantilla(), self.mantenimiento, hasta=S3, aceptar=True)
        self.assertEqual(
            set(res.serie.reservas.values_list("estado", flat=True)), {Estado.ACEPTADA}
        )
        self.assertEqual(len(mail.outbox), 0)
        for r in res.serie.reservas.all():
            self.assertEqual(r.gestionada_por, self.mantenimiento)
            self.assertEqual([e.tipo for e in r.eventos.all()], [Tipo.CREADA, Tipo.ACEPTADA])

    def test_omite_las_fechas_ocupadas_y_crea_las_demas(self):
        self.crear(dia=S2, ini="11:00", fin="12:00", estado=Estado.ACEPTADA)  # solapa 10:00-12:20
        res = crear_serie(self.plantilla(), self.profesor, hasta=S4)
        self.assertEqual(res.omitidas, [S2])
        self.assertEqual([r.fecha for r in res.creadas], [S1, S3, S4])

    def test_tambien_omite_al_reservar_directamente(self):
        self.crear(dia=S3, estado=Estado.ACEPTADA)
        res = crear_serie(self.plantilla(), self.mantenimiento, hasta=S4, aceptar=True)
        self.assertEqual(res.omitidas, [S3])
        self.assertEqual(len(res.creadas), 3)

    def test_si_todo_esta_ocupado_no_se_guarda_nada(self):
        for dia in (S1, S2):
            self.crear(dia=dia, ini="10:00", fin="13:00", estado=Estado.ACEPTADA)
        antes = Reserva.objects.count()
        with self.assertRaises(ConflictoReserva):
            crear_serie(self.plantilla(), self.profesor, hasta=S2)
        self.assertEqual((Reserva.objects.count(), SerieReserva.objects.count()), (antes, 0))

    def test_rechaza_series_mas_largas_que_el_maximo(self):
        with self.assertRaises(ValueError):
            crear_serie(self.plantilla(), self.profesor, hasta=S1 + timedelta(days=7 * 40))
        self.assertEqual(SerieReserva.objects.count(), 0)


class SerieDecisionTests(BaseTest):
    def setUp(self):
        plantilla = Reserva(
            aula=self.a1, fecha=S1, hora_inicio=t("10:00"), hora_fin=t("12:20"),
            profesor="Ana López", titulo="Anatomía I", curso=self.curso,
        )
        self.serie = crear_serie(plantilla, self.profesor, hasta=S4).serie

    def estados(self):
        return list(self.serie.reservas.order_by("fecha").values_list("estado", flat=True))

    def test_aceptar_la_serie_acepta_todas_con_auditoria_individual(self):
        res = cambiar_estado_serie(self.serie, Estado.ACEPTADA, self.mantenimiento)
        self.assertEqual((len(res.cambiadas), res.en_conflicto), (4, []))
        self.assertEqual(self.estados(), [Estado.ACEPTADA] * 4)
        for r in self.serie.reservas.all():
            self.assertEqual((r.gestionada_por, r.eventos.last().tipo), (self.mantenimiento, Tipo.ACEPTADA))
            self.assertIsNotNone(r.gestionada_en)

    def test_acepta_las_libres_y_deja_pendientes_las_ocupadas(self):
        # Otra reserva se acepta en el aula después de crear la serie.
        otra = self.crear(dia=S2, ini="09:00", fin="11:00", estado=Estado.ACEPTADA)
        res = cambiar_estado_serie(self.serie, Estado.ACEPTADA, self.mantenimiento)
        self.assertEqual([r.fecha for r in res.en_conflicto], [S2])
        self.assertEqual(
            self.estados(), [Estado.ACEPTADA, Estado.PENDIENTE, Estado.ACEPTADA, Estado.ACEPTADA]
        )
        self.assertEqual(otra.estado, Estado.ACEPTADA)

    def test_rechazar_la_serie_guarda_el_motivo_en_cada_fecha(self):
        res = cambiar_estado_serie(self.serie, Estado.RECHAZADA, self.mantenimiento, "Aula en obras")
        self.assertEqual(len(res.cambiadas), 4)
        for r in self.serie.reservas.all():
            self.assertEqual((r.estado, r.motivo_rechazo), (Estado.RECHAZADA, "Aula en obras"))
            self.assertEqual(r.eventos.last().detalle, "Aula en obras")

    def test_no_toca_las_fechas_ya_decididas(self):
        s1, s2, *_ = self.serie.reservas.order_by("fecha")
        cambiar_estado(s1, Estado.ACEPTADA, self.mantenimiento)
        cambiar_estado(s2, Estado.RECHAZADA, self.informatica, "Festivo")
        res = cambiar_estado_serie(self.serie, Estado.RECHAZADA, self.mantenimiento, "Se cancela el resto")
        self.assertEqual(len(res.cambiadas), 2)
        s1.refresh_from_db(), s2.refresh_from_db()
        self.assertEqual((s1.estado, s2.estado, s2.motivo_rechazo), (Estado.ACEPTADA, Estado.RECHAZADA, "Festivo"))

    def test_repetir_la_accion_no_duplica_el_historial(self):
        cambiar_estado_serie(self.serie, Estado.ACEPTADA, self.mantenimiento)
        eventos = EventoReserva.objects.count()
        res = cambiar_estado_serie(self.serie, Estado.ACEPTADA, self.mantenimiento)
        self.assertEqual((res.cambiadas, res.en_conflicto), ([], []))
        self.assertEqual(EventoReserva.objects.count(), eventos)

    def test_una_serie_solo_se_acepta_o_rechaza(self):
        with self.assertRaises(ValueError):
            cambiar_estado_serie(self.serie, Estado.PENDIENTE, self.mantenimiento)

    def test_la_restriccion_de_la_base_de_datos_sigue_protegiendo_las_series(self):
        cambiar_estado_serie(self.serie, Estado.ACEPTADA, self.mantenimiento)
        with self.assertRaises(IntegrityError), transaction.atomic():
            Reserva.objects.create(
                aula=self.a1, fecha=S3, hora_inicio=t("11:00"), hora_fin=t("12:00"),
                profesor="X", estado=Estado.ACEPTADA, creada_por=self.profesor,
            )


class SerieVistasTests(BaseTest):
    def pedir(self, usuario=None, **cambios):
        self.client.force_login(usuario or self.profesor)
        datos = self.datos_form(
            **{"fecha": S1.isoformat(), "hora_fin": "12:20", "repetir_hasta": S4.isoformat(), **cambios}
        )
        with self.captureOnCommitCallbacks(execute=True):
            return self.client.post(reverse("reservar"), datos, follow=True)

    def test_el_profesor_pide_una_serie(self):
        r = self.pedir()
        self.assertRedirects(r, reverse("reservar"))
        self.assertEqual(Reserva.objects.count(), 4)
        self.assertEqual(Reserva.objects.values("serie").distinct().count(), 1)
        self.assertContains(r, "Petición enviada: 4 fechas")
        self.assertContains(r, "10:00 a 12:20")
        self.assertEqual(len(mail.outbox), 1)

    def test_avisa_de_las_fechas_que_no_se_han_incluido(self):
        self.crear(dia=S2, estado=Estado.ACEPTADA)
        r = self.pedir()
        self.assertEqual(Reserva.objects.filter(serie__isnull=False).count(), 3)
        self.assertContains(r, "No se han incluido 1 fechas")
        self.assertContains(r, f"{S2.day}/{S2.month}")

    def test_el_admin_reserva_una_serie_directamente(self):
        r = self.pedir(self.mantenimiento, aceptar_directamente="on")
        self.assertContains(r, "Serie creada y aceptada: 4 fechas")
        self.assertEqual(Reserva.objects.filter(estado=Estado.ACEPTADA).count(), 4)
        self.assertEqual(len(mail.outbox), 0)

    def test_un_dato_invalido_no_crea_nada(self):
        r = self.pedir(repetir_hasta=S1.isoformat())
        self.assertEqual(Reserva.objects.count(), 0)
        self.assertContains(r, "al menos una semana posterior")

    def test_sin_repetir_sigue_creando_una_sola_reserva(self):
        self.client.force_login(self.profesor)
        self.client.post(reverse("reservar"), self.datos_form())
        reserva = Reserva.objects.get()
        self.assertIsNone(reserva.serie)

    def test_el_formulario_muestra_el_campo(self):
        self.client.force_login(self.profesor)
        self.assertContains(self.client.get(reverse("reservar")), 'name="repetir_hasta"')

    def test_aviso_en_vivo_de_la_serie(self):
        self.crear(dia=S2, estado=Estado.ACEPTADA)
        self.client.force_login(self.profesor)
        params = {
            "aula": self.a1.pk, "fecha": S1.isoformat(), "hora_inicio": "10:00",
            "hora_fin": "11:00", "repetir_hasta": S4.isoformat(),
        }
        r = self.client.get(reverse("disponibilidad"), params)
        self.assertContains(r, "1 de 4 fechas con el aula ocupada")
        self.assertContains(r, "Se pedirán solo las otras 3")
        libre = self.client.get(reverse("disponibilidad"), {**params, "aula": self.a2.pk})
        self.assertContains(libre, "Las 4 fechas están libres")

    def test_aviso_en_vivo_ignora_una_fecha_final_invalida(self):
        self.client.force_login(self.profesor)
        params = {
            "aula": self.a1.pk, "fecha": S1.isoformat(), "hora_inicio": "10:00",
            "hora_fin": "11:00", "repetir_hasta": "basura",
        }
        self.assertContains(self.client.get(reverse("disponibilidad"), params), "Franja libre")

    # ------------------------------------------------------------------ panel y detalle
    def serie_pendiente(self):
        self.pedir()
        return Reserva.objects.order_by("fecha").first()

    def test_la_cola_muestra_la_serie_como_un_solo_elemento(self):
        self.serie_pendiente()
        self.crear(dia=S1, aula=self.a2)  # una petición suelta
        self.client.force_login(self.mantenimiento)
        r = self.client.get(reverse("panel_contenido"), {"semana": S1.isoformat()})
        self.assertContains(r, "Aceptar serie", count=1)
        self.assertContains(r, "4 fechas, del")
        self.assertEqual(r.context["total_pendientes"], 2)
        self.assertEqual(len(r.context["pendientes"]), 2)

    def test_la_cola_avisa_de_fechas_ocupadas_de_la_serie(self):
        self.serie_pendiente()
        self.crear(dia=S3, estado=Estado.ACEPTADA, ini="11:00", fin="12:00", aula=self.a1)
        self.client.force_login(self.mantenimiento)
        r = self.client.get(reverse("panel_contenido"), {"semana": S1.isoformat()})
        self.assertContains(r, "1 de 4 fechas ocupadas")
        self.assertContains(r, "Aceptar serie")  # se puede aceptar: las libres entran

    def test_la_cola_no_deja_aceptar_si_todas_las_fechas_estan_ocupadas(self):
        self.serie_pendiente()
        for dia in (S1, S2, S3, S4):
            self.crear(dia=dia, estado=Estado.ACEPTADA, ini="10:00", fin="13:00")
        self.client.force_login(self.mantenimiento)
        r = self.client.get(reverse("panel_contenido"), {"semana": S1.isoformat()})
        self.assertNotContains(r, "Aceptar serie")
        self.assertContains(r, "Aula ya ocupada")

    def test_el_detalle_del_admin_muestra_la_serie(self):
        primera = self.serie_pendiente()
        self.client.force_login(self.mantenimiento)
        r = self.client.get(reverse("reserva_detalle", args=[primera.pk]), HTTP_HX_REQUEST="true")
        self.assertContains(r, "Serie semanal (4 fechas)")
        self.assertContains(r, "Aceptar las 4 pendientes")
        self.assertContains(r, "Rechazar las 4 pendientes")
        self.assertContains(r, reverse("reserva_serie_accion", args=[primera.pk]))

    def test_una_reserva_suelta_no_muestra_bloque_de_serie(self):
        suelta = self.crear()
        self.client.force_login(self.mantenimiento)
        r = self.client.get(reverse("reserva_detalle", args=[suelta.pk]), HTTP_HX_REQUEST="true")
        self.assertNotContains(r, "Serie semanal")

    def test_el_profesorado_no_ve_series_pendientes_ni_sus_acciones(self):
        primera = self.serie_pendiente()
        self.client.force_login(self.profesor)
        r = self.client.get(reverse("panel_contenido"), {"semana": S1.isoformat()})
        self.assertNotContains(r, "Anatomía I")
        self.assertNotContains(r, "Aceptar serie")
        self.assertEqual(
            self.client.get(reverse("reserva_detalle", args=[primera.pk]), HTTP_HX_REQUEST="true").status_code, 404
        )

    # ------------------------------------------------------------------ acciones sobre la serie
    def test_aceptar_la_serie_desde_el_panel(self):
        primera = self.serie_pendiente()
        self.client.force_login(self.mantenimiento)
        r = self.client.post(
            reverse("reserva_serie_accion", args=[primera.pk]), {"accion": "aceptar"}, HTTP_HX_REQUEST="true"
        )
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "4 fechas aceptadas.")
        self.assertEqual(r["HX-Trigger"], "reservasActualizadas")
        self.assertEqual(Reserva.objects.filter(estado=Estado.ACEPTADA).count(), 4)

    def test_rechazar_la_serie_con_motivo(self):
        primera = self.serie_pendiente()
        self.client.force_login(self.mantenimiento)
        r = self.client.post(
            reverse("reserva_serie_accion", args=[primera.pk]),
            {"accion": "rechazar", "motivo": "Obras"}, HTTP_HX_REQUEST="true",
        )
        self.assertContains(r, "4 fechas rechazadas. Recuerda avisar al profesor.")
        self.assertEqual(set(Reserva.objects.values_list("motivo_rechazo", flat=True)), {"Obras"})

    def test_aceptacion_parcial_informa_de_las_fechas_pendientes(self):
        primera = self.serie_pendiente()
        self.crear(dia=S3, estado=Estado.ACEPTADA, ini="11:00", fin="12:00")
        self.client.force_login(self.mantenimiento)
        r = self.client.post(
            reverse("reserva_serie_accion", args=[primera.pk]), {"accion": "aceptar"}, HTTP_HX_REQUEST="true"
        )
        self.assertContains(r, "3 fechas aceptadas.")
        self.assertContains(r, "1 fechas siguen pendientes")
        self.assertContains(r, f"{S3.day}/{S3.month}")
        self.assertEqual(Reserva.objects.filter(serie__isnull=False, estado=Estado.PENDIENTE).count(), 1)

    def test_si_ninguna_se_puede_aceptar_se_marca_como_error(self):
        primera = self.serie_pendiente()
        for dia in (S1, S2, S3, S4):
            self.crear(dia=dia, estado=Estado.ACEPTADA, ini="10:00", fin="13:00")
        self.client.force_login(self.mantenimiento)
        r = self.client.post(
            reverse("reserva_serie_accion", args=[primera.pk]), {"accion": "aceptar"}, HTTP_HX_REQUEST="true"
        )
        self.assertContains(r, "aviso-error")
        self.assertContains(r, "0 fechas aceptadas.")

    def test_sin_htmx_redirige_y_muestra_el_mensaje(self):
        primera = self.serie_pendiente()
        self.client.force_login(self.mantenimiento)
        r = self.client.post(reverse("reserva_serie_accion", args=[primera.pk]), {"accion": "aceptar"}, follow=True)
        self.assertRedirects(r, reverse("reserva_detalle", args=[primera.pk]))
        self.assertContains(r, "4 fechas aceptadas.")

    def test_acciones_no_validas(self):
        primera = self.serie_pendiente()
        suelta = self.crear(dia=S1, aula=self.a2)
        self.client.force_login(self.mantenimiento)
        url = reverse("reserva_serie_accion", args=[primera.pk])
        self.assertEqual(self.client.post(url, {"accion": "reabrir"}).status_code, 400)
        self.assertEqual(self.client.post(url, {}).status_code, 400)
        self.assertEqual(
            self.client.post(reverse("reserva_serie_accion", args=[suelta.pk]), {"accion": "aceptar"}).status_code, 400
        )
        self.assertEqual(self.client.get(url).status_code, 405)

    def test_el_profesorado_no_puede_decidir_una_serie(self):
        primera = self.serie_pendiente()
        self.client.force_login(self.profesor)
        r = self.client.post(reverse("reserva_serie_accion", args=[primera.pk]), {"accion": "aceptar"})
        self.assertEqual(r.status_code, 403)
        self.assertEqual(Reserva.objects.filter(estado=Estado.ACEPTADA).count(), 0)

    def test_anonimo_va_al_login(self):
        r = self.client.post(reverse("reserva_serie_accion", args=[1]), {"accion": "aceptar"})
        self.assertEqual(r.status_code, 302)
        self.assertIn("login", r["Location"])

    def test_una_fecha_se_puede_decidir_aparte_de_la_serie(self):
        primera = self.serie_pendiente()
        self.client.force_login(self.mantenimiento)
        self.client.post(reverse("reserva_accion", args=[primera.pk]), {"accion": "rechazar", "motivo": "Festivo"})
        r = self.client.post(
            reverse("reserva_serie_accion", args=[primera.pk]), {"accion": "aceptar"}, HTTP_HX_REQUEST="true"
        )
        self.assertContains(r, "3 fechas aceptadas.")
        primera.refresh_from_db()
        self.assertEqual(primera.estado, Estado.RECHAZADA)

    def test_exportacion_incluye_la_serie(self):
        self.serie_pendiente()
        self.crear(dia=S1, aula=self.a2)
        self.client.force_login(self.mantenimiento)
        r = self.client.get(reverse("exportar"), {"formato": "reservas_csv"})
        filas = list(csv.DictReader(StringIO(r.content.decode("utf-8-sig"))))
        series = [f["serie"] for f in filas]
        self.assertEqual(len([x for x in series if x]), 4)
        self.assertEqual(len(set(x for x in series if x)), 1)
        self.assertEqual(series.count(""), 1)

    def test_el_admin_de_django_no_permite_cambiar_la_serie(self):
        primera = self.serie_pendiente()
        self.client.force_login(self.informatica)
        r = self.client.get(reverse("admin:reservas_reserva_change", args=[primera.pk]))
        self.assertEqual(r.status_code, 200)
        self.assertNotContains(r, 'name="serie"')